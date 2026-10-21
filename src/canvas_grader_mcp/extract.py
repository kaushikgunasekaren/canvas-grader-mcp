"""Turn assignment descriptions and submitted files into plain text for grading."""

from __future__ import annotations

import json
import re
from html.parser import HTMLParser
from pathlib import Path

MAX_CHARS = 60_000

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".py", ".java", ".c", ".h", ".cpp", ".hpp", ".cc", ".cs", ".js", ".jsx", ".ts", ".tsx", ".go", ".rs",
    ".rb", ".php", ".swift", ".kt", ".kts", ".scala", ".m", ".r", ".sql", ".sh", ".bash", ".ps1",
    ".tex", ".bib", ".css", ".scss", ".asm", ".s", ".v", ".sv", ".vhd", ".hs", ".ml", ".lua", ".pl", ".dart",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
MEDIA_EXTENSIONS = {".mp4", ".mov", ".m4v", ".avi", ".mkv", ".webm", ".mp3", ".m4a", ".wav"}
MAX_MEMBER_BYTES = 100 * 1024 * 1024


def unzip(path: Path) -> list[Path]:
    """Safely extract a submitted .zip next to it; media files are listed but not extracted."""
    import zipfile

    dest = path.with_name(path.name + ".d")
    out = []
    with zipfile.ZipFile(path) as zf:
        for info in zf.infolist():
            name = Path(info.filename)
            if info.is_dir() or name.is_absolute() or ".." in name.parts or name.parts[0] == "__MACOSX":
                continue
            target = dest / name
            if name.suffix.lower() in MEDIA_EXTENSIONS or info.file_size > MAX_MEMBER_BYTES:
                out.append(target)  # recorded, not extracted
                continue
            if not target.exists():
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(zf.read(info))
            out.append(target)
    return out


class _HTMLToText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote", "table", "ul", "ol"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0
        self._href: str | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")
            if tag == "li":
                self.parts.append("- ")
            elif tag.startswith("h") and len(tag) == 2:
                self.parts.append("#" * int(tag[1]) + " ")
        elif tag == "a":
            self._href = dict(attrs).get("href")
        elif tag == "img":
            alt = dict(attrs).get("alt")
            self.parts.append(f"[image{': ' + alt if alt else ''}]")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif tag in self.BLOCK:
            self.parts.append("\n")
        elif tag == "a" and self._href:
            self.parts.append(f" ({self._href})")
            self._href = None

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def html_to_text(html: str | None) -> str:
    if not html:
        return ""
    parser = _HTMLToText()
    parser.feed(html)
    text = "".join(parser.parts)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"^([ \t]*-)[ \t]*\n+[ \t]*", r"\1 ", text, flags=re.M)  # <li><p>..</p></li>
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def truncate(text: str, limit: int = MAX_CHARS) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n[... truncated: {len(text) - limit} more characters not shown ...]"


def extract_file(path: Path) -> dict:
    """Return {'kind', 'text'?, 'note'?} for a downloaded submission file."""
    from .media import effective_suffix, video_duration

    ext = effective_suffix(path)
    detected = {"detected_type": ext} if ext != path.suffix.lower() else {}
    result = _extract(path, ext)
    if result.get("kind") == "media" and path.exists():
        if duration := video_duration(path):
            result["duration_seconds"] = round(duration, 1)
            result["note"] += " Use get_submission_images(video_frames=N) to see frames (e.g. a recording of code)."
    return {**result, **detected}


def _extract(path: Path, ext: str) -> dict:
    try:
        if ext == ".pdf":
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(path))
            pages, image_pages = [], []
            for i, page in enumerate(pdf):
                text = page.get_textpage().get_text_range().replace("\r\n", "\n").strip()
                images = count_pdf_images(page)
                marker = f" [{images} image(s) - see get_submission_images]" if images else ""
                if images:
                    image_pages.append(i + 1)
                links = pdf_links(pdf, page)
                if links:  # e.g. a video link hidden behind "click here"
                    text += "\n[hyperlinks on this page: " + ", ".join(links) + "]"
                pages.append(f"--- page {i + 1}{marker} ---\n{text}")
            text = "\n\n".join(pages)
            has_text = any(p.split("---\n", 1)[-1].strip() for p in pages)
            note = None if has_text else "No extractable text (scanned PDF?); use get_submission_images."
            return {"kind": "pdf", "pages": len(pdf), "pages_with_images": image_pages, "text": truncate(text), "note": note}
        if ext == ".docx":
            import docx

            doc = docx.Document(str(path))
            images = sum(1 for rel in doc.part.rels.values() if "image" in rel.reltype)
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            note = f"Contains {images} embedded image(s); see get_submission_images." if images else None
            return {"kind": "docx", "images": images, "text": truncate("\n".join(lines)), "note": note}
        if ext == ".ipynb":
            nb = json.loads(path.read_text(errors="replace"))
            cells = []
            for i, cell in enumerate(nb.get("cells", [])):
                src = "".join(cell.get("source", []))
                cells.append(f"--- cell {i + 1} ({cell.get('cell_type')}) ---\n{src}")
                for out in cell.get("outputs", []):
                    out_text = "".join(out.get("text", [])) or "".join(out.get("data", {}).get("text/plain", []))
                    if out_text:
                        cells.append(f"[output]\n{out_text}")
            return {"kind": "notebook", "text": truncate("\n\n".join(cells))}
        if ext in (".html", ".htm"):
            return {"kind": "html", "text": truncate(html_to_text(path.read_text(errors="replace")))}
        if ext in TEXT_EXTENSIONS or _looks_like_text(path):
            return {"kind": "text", "text": truncate(path.read_text(errors="replace"))}
        if ext == ".sb3":
            return {"kind": "blocks project", "text": truncate(describe_sb3(path))}
        if ext == ".zip":
            import zipfile

            sizes = {Path(i.filename).name: i.file_size for i in zipfile.ZipFile(path).infolist()}
            members = []
            for member in unzip(path):
                entry = {"filename": str(member.relative_to(path.with_name(path.name + ".d"))), "size": sizes.get(member.name)}
                if member.suffix.lower() in MEDIA_EXTENSIONS:
                    entry.update(kind="media", note="Video/audio file inside the zip (not extracted, can't be watched here).")
                elif not member.exists():
                    entry.update(kind="skipped", note="Too large to extract.")
                else:
                    entry.update(extract_file(member))
                members.append(entry)
            return {"kind": "zip", "members": members}
        if ext in MEDIA_EXTENSIONS:
            return {"kind": "media", "note": "Video/audio file: it was submitted, but can't be watched here."}
        if ext in IMAGE_EXTENSIONS:
            return {"kind": "image", "note": "Image file; see get_submission_images."}
    except Exception as exc:  # a corrupt upload should not break grading the rest
        return {"kind": "error", "note": f"Could not read {path.name}: {exc}"}
    return {"kind": "binary", "note": f"Unsupported file type {ext or '(none)'}; review it in SpeedGrader."}


def describe_sb3(path: Path) -> str:
    """Readable listing of a Scratch 3 project (.sb3, also used by Scratch-based tools): each script, nested by C-blocks."""
    import zipfile

    project = json.loads(zipfile.ZipFile(path).read("project.json"))
    out: list[str] = []
    for target in project.get("targets", []):
        blocks = {k: b for k, b in (target.get("blocks") or {}).items() if isinstance(b, dict)}

        def value(inp) -> str:
            # inputs are [shadow_kind, literal-or-block-id, ...]; literals are [type, value]
            for part in inp[1:]:
                if isinstance(part, list) and len(part) >= 2:
                    return str(part[1])
                if isinstance(part, str) and part in blocks:
                    return f"({label(part)})"
            return "?"

        def label(bid: str) -> str:
            b = blocks[bid]
            args = [f"{k}={v[0]}" for k, v in (b.get("fields") or {}).items()]
            args += [f"{k}={value(v)}" for k, v in (b.get("inputs") or {}).items() if not k.startswith("SUBSTACK")]
            return b.get("opcode", "?") + (" " + " ".join(args) if args else "")

        def walk(bid: str | None, depth: int) -> None:
            while bid and bid in blocks:
                out.append("  " * depth + label(bid))
                for key in ("SUBSTACK", "SUBSTACK2"):
                    sub = (blocks[bid].get("inputs") or {}).get(key)
                    if sub and len(sub) > 1 and isinstance(sub[1], str):
                        out.append("  " * (depth + 1) + ("else:" if key == "SUBSTACK2" else "do:"))
                        walk(sub[1], depth + 2)
                bid = blocks[bid].get("next")

        tops = [k for k, b in blocks.items() if b.get("topLevel") and not b.get("shadow")]
        for i, top in enumerate(tops, 1):
            out.append(f"--- {target.get('name', 'sprite')} script {i} ---")
            walk(top, 0)
    return "\n".join(out) or "(empty project)"


def pdf_links(pdf, page) -> list[str]:
    """URLs of clickable links on a PDF page."""
    import ctypes

    import pypdfium2.raw as raw

    urls: list[str] = []
    pos, link = ctypes.c_int(0), raw.FPDF_LINK()
    while raw.FPDFLink_Enumerate(page.raw, ctypes.byref(pos), ctypes.byref(link)):
        action = raw.FPDFLink_GetAction(link)
        if not action or raw.FPDFAction_GetType(action) != raw.PDFACTION_URI:
            continue
        size = raw.FPDFAction_GetURIPath(pdf.raw, action, None, 0)
        buf = ctypes.create_string_buffer(size)
        raw.FPDFAction_GetURIPath(pdf.raw, action, buf, size)
        url = buf.value.decode(errors="replace")
        if url and url not in urls:
            urls.append(url)
    return urls


def count_pdf_images(page) -> int:
    import pypdfium2 as pdfium

    return sum(1 for obj in page.get_objects(max_depth=2) if obj.type == pdfium.raw.FPDF_PAGEOBJ_IMAGE)


def _looks_like_text(path: Path) -> bool:
    try:
        chunk = path.read_bytes()[:4096]
    except OSError:
        return False
    if b"\x00" in chunk:
        return False
    try:
        chunk.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return True
